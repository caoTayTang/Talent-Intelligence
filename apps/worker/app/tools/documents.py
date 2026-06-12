from io import BytesIO

import pdfplumber
from docx import Document


def parse_pdf_bytes(file_bytes: bytes) -> str:
    """
    Extract text and table content from a PDF
    Tables are converted into line-based text so the scorer can use them
    """
    chunks: list[str] = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(f"\n[Page {page_idx} Text]\n{text}")

            tables = page.extract_tables() or []
            for tbl_idx, table in enumerate(tables, start=1):
                chunks.append(f"\n[Page {page_idx} Table {tbl_idx}]")
                for row in table:
                    clean_row = [cell.strip() if cell else "" for cell in row]
                    chunks.append(" | ".join(clean_row))
    return "\n".join(chunks).strip()


def parse_docx_bytes(file_bytes: bytes) -> str:
    """
    Extract paragraphs and tables from a Docx file
    """
    document = Document(BytesIO(file_bytes))
    chunks: list[str] = []
    for paragaph in document.paragraphs:
        text = paragaph.text.strip()
        if text:
            chunks.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"\n[Table {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            chunks.append(" | ".join(cells))

    return "\n".join(chunks).strip()


def parse_document_bytes(file_bytes: bytes, object_key: str) -> str:
    """
    Parse a CV object based on the file ext
    """
    lower_key = object_key.lower()
    if lower_key.endswith(".pdf"):
        return parse_pdf_bytes(file_bytes)
    elif lower_key.endswith(".docx"):
        return parse_docx_bytes(file_bytes)

    return file_bytes.decode("utf-8", errors="ignore")
